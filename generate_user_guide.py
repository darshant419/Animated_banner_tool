#!/usr/bin/env python3
"""
Generate the Banner Tool - Elements, Usage & Functionality Guide (Word format).

Covers every UI element in the tool, how to use it, and what it does,
derived from the actual source code (src/ components, store, utils, templates).
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUTPUT = "Banner_Tool_Elements_Usage_Functionality_Guide.docx"


# ----------------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------------
def shade_cell(cell, color):
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color)
    cell._element.get_or_add_tcPr().append(shading)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ''
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade_cell(hdr[i], '2E2E3D')
    for r_i, row in enumerate(rows):
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ''
            p = cells[i].paragraphs[0]
            run = p.add_run(str(val))
            run.font.size = Pt(9)
            if r_i % 2 == 1:
                shade_cell(cells[i], 'F5F6F8')
    if widths:
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return table


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style='List Bullet')


def add_numbered(doc, items):
    for item in items:
        doc.add_paragraph(item, style='List Number')


def add_kv(doc, pairs):
    """Bold key + description paragraph pairs."""
    for key, desc in pairs:
        p = doc.add_paragraph()
        r = p.add_run(f'{key}: ')
        r.bold = True
        p.add_run(desc)


# ==== APPEND MARK ====


# ----------------------------------------------------------------------------
# Document
# ----------------------------------------------------------------------------
def build():
    doc = Document()

    # Base style
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10.5)

    # ---- Title page ----
    title = doc.add_heading('Banner Tool', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_heading('Elements, Usage & Functionality Guide', level=1)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph('Complete reference for every tool, panel, element, control, '
                          'keyboard shortcut, and workflow in the application.')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True
    p.runs[0].font.size = Pt(12)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run('React + TypeScript + Konva + GSAP + Zustand\n').bold = True
    meta.add_run('Generated from source code · September 2026')

    doc.add_page_break()

    # ---- Table of contents (manual) ----
    doc.add_heading('Table of Contents', level=1)
    toc = [
        '1. Application Overview',
        '2. Getting Started — Workflow Modes',
        '3. Screen Layout',
        '4. Header Bar Controls',
        '5. Left Toolbar — Tools & Usage',
        '6. Designable Element Types',
        '7. Canvas (Design Surface)',
        '8. Properties Panel (Right Side)',
        '9. Left Panels: Layers, Assets, Templates, Sizes',
        '10. Timeline (Bottom Panel)',
        '11. Animation Studio (Animation Dialog)',
        '12. ISI Pro Suite (Compliance Scroll)',
        '13. Artboards & Multi-size Campaigns',
        '14. Export Functionality',
        '15. Templates Catalog',
        '16. Keyboard Shortcuts',
        '17. State Management & Undo/Redo',
        '18. Technology Stack',
    ]
    for item in toc:
        doc.add_paragraph(item, style='List Bullet')
    doc.add_page_break()

    # =====================================================================
    # 1. Overview
    # =====================================================================
    doc.add_heading('1. Application Overview', level=1)
    doc.add_paragraph(
        'The Banner Tool is a professional web-based design application for creating static and '
        'animated HTML banners, purpose-built for pharmaceutical and healthcare marketing. It '
        'provides a visual canvas editor, a keyframe/preset animation system, a dedicated ISI '
        '(Important Safety Information) compliance component, multi-size (IAB) artboard '
        'management, and one-click export to PNG or a production-ready HTML/CSS/JS banner package.'
    )
    doc.add_heading('Key Capabilities', level=2)
    add_bullets(doc, [
        'Two workflow modes: static EMR banners and animated multi-frame sequences.',
        'Eight designable element types: text, rectangle, circle, image, video, SVG shape, ISI scroll, and raw HTML.',
        '300+ preset animations (Animista-inspired catalog) grouped into Entrance, Emphasis/Loop, and Exit, plus hover effects.',
        'Frame-accurate timeline with manual keyframes, easing curves, scrubbing, and stagger sequencing.',
        'Multi-artboard support with smart reflow and live multi-size previews of 10 standard IAB sizes.',
        'ISI Pro Suite: scrollable safety-information tray with auto-scroll, header bar, logo, scrollbar and typography controls.',
        'Full undo/redo history with state snapshots.',
        'Export to PNG (2x resolution) and to a ZIP package (index.html + css/styles.css + js/main.js + images/).',
        'Pre-built medical templates with pre-loaded ISI content.',
    ])

    # =====================================================================
    # 2. Getting started
    # =====================================================================
    doc.add_heading('2. Getting Started — Workflow Modes', level=1)
    doc.add_paragraph(
        'On launch the Mode Select screen asks you to pick a starting point. You can switch '
        'modes at any time from the header ("Switch mode" button), which returns you to this screen.'
    )
    add_table(
        doc,
        ['Mode', 'Card Label', 'What It Does'],
        [
            ['EMR Banner (Static)', 'EMR · Static banners · All sizes',
             'Starts from the EMR standard layout with a pre-loaded ISI scroll area. Suited to static medical banners.'],
            ['Animated Banner', 'Animated · Frames / Animations / ISI',
             'Starts an animated multi-frame sequence with keyframes, animation presets, and ISI, ready to customize.'],
        ],
        widths=[1.4, 2.0, 3.1],
    )
    doc.add_paragraph('Both modes open the same editor; the mode seeds the starting template/state.')

    # =====================================================================
    # 3. Screen layout
    # =====================================================================
    doc.add_heading('3. Screen Layout', level=1)
    doc.add_paragraph('The editor (MainLayout) is arranged as follows:')
    add_table(
        doc,
        ['Region', 'Location', 'Component', 'Purpose'],
        [
            ['Header bar', 'Top', 'MainLayout header', 'Mode switch, artboard tabs, multi-size, undo/redo, Preview, Export HTML, PNG'],
            ['Toolbar', 'Far left (64 px)', 'Toolbar.tsx', 'Tool selection: Select, Text, Image, ISI, Assets, Templates, Layers, Sizes'],
            ['Side panel', 'Left of canvas', 'Layers / Assets / Templates / Variations panels', 'Opened by the matching toolbar tool'],
            ['Canvas', 'Center', 'DesignCanvas.tsx (Konva stage)', 'Visual design surface; single or multi-artboard view'],
            ['Timeline', 'Bottom', 'Timeline.tsx', 'Playhead, layer tracks, animation segments, keyframes, playback'],
            ['Properties', 'Right (320 px)', 'PropertiesPanel.tsx', 'Canvas settings when nothing selected; element/keyframe properties when selected'],
            ['Modals', 'Overlay', 'AnimationDialog, Multi-size campaign', 'Animation Studio and batch size generation'],
        ],
        widths=[1.1, 1.1, 1.7, 2.6],
    )

    # =====================================================================
    # 4. Header bar
    # =====================================================================
    doc.add_heading('4. Header Bar Controls', level=1)
    add_table(
        doc,
        ['Control', 'Usage', 'Functionality'],
        [
            ['Mode badge / Switch mode', 'Click "Switch mode" (refresh icon)',
             'Returns to the Mode Select screen to pick EMR or Animated again.'],
            ['Artboard tabs', 'Click a size chip (e.g. 300×250)',
             'Switches the active artboard; each artboard keeps its own element layout.'],
            ['+ (Add size)', 'Click "+" with chevron',
             'Dropdown of IAB presets; adds a new artboard of the chosen size.'],
            ['Multi-size', 'Click "Multi-size"',
             'Opens the Multi-size campaign modal to batch-generate sizes from the current design.'],
            ['Undo', 'Click or Ctrl+Z', 'Reverts the last state snapshot (disabled when history empty).'],
            ['Redo', 'Click or Ctrl+Y (or Ctrl+Shift+Z)', 'Re-applies a reverted snapshot.'],
            ['Preview / Stop', 'Click the red Preview button',
             'Toggles GSAP playback of the whole timeline (Play to Pause/Stop).'],
            ['Export HTML', 'Click "Export HTML"',
             'Builds and downloads a ZIP banner package (index.html, css, js, images).'],
            ['PNG', 'Click "PNG"',
             'Downloads the active artboard as a PNG at 2x pixel ratio, named banner-WxH.png.'],
        ],
        widths=[1.3, 1.9, 3.3],
    )

    # =====================================================================
    # 5. Toolbar
    # =====================================================================
    doc.add_heading('5. Left Toolbar — Tools & Usage', level=1)
    doc.add_paragraph(
        'The vertical toolbar is the primary way to add content. Tools that create elements '
        'insert a default instance on the canvas and automatically return to the Select tool; '
        'panel tools toggle the corresponding left side panel.'
    )
    add_table(
        doc,
        ['Tool (Label)', 'Icon', 'Usage', 'Functionality'],
        [
            ['Select', 'Mouse pointer', 'Click the tool (or after any creation tool)',
             'Default mode: click elements to select, drag to move, use transformer handles to resize/rotate.'],
            ['Text', 'Type icon', 'Click once',
             'Adds a text element (default "Double click to edit", 24 px) at 100,80 sized 200x60; returns to Select. Double-click on canvas edits text inline.'],
            ['Image', 'Image icon', 'Click once',
             'Opens a hidden file picker (accept="image/*"). The chosen file is read as a data URL and inserted at 100,100 with aspect-ratio-preserving height (300 px wide).'],
            ['ISI', 'Scroll-text icon', 'Click once',
             'Inserts an ISI Scroll element (300x88 at y=162) pre-loaded with the EMR ISI text, scroll speed 30, auto-start on, 12 px green type; returns to Select.'],
            ['Assets', 'Shapes icon', 'Click to open panel',
             'Toggles the Assets panel (upload/search library).'],
            ['Templates', 'File-JSON icon', 'Click to open panel',
             'Toggles the Template Gallery panel.'],
            ['Layers', 'Layers icon', 'Click to open panel',
             'Toggles the Layers panel.'],
            ['Sizes', 'Layout icon', 'Click to open panel',
             'Toggles the Multi-size Preview (Variations) panel.'],
        ],
        widths=[1.0, 0.9, 1.5, 3.1],
    )

    # =====================================================================
    # 6. Element types
    # =====================================================================
    doc.add_heading('6. Designable Element Types', level=1)
    doc.add_paragraph('The store defines eight kinds of objects (ElementType). All are rendered '
                      'on the Konva canvas and share common properties (position, size, rotation, '
                      'opacity, visibility, lock, name, animation).')
    add_table(
        doc,
        ['Type', 'How Added', 'Description & Notable Properties'],
        [
            ['text', 'Toolbar to Text',
             'Typography with content, font size/family/weight/style, alignment (left/center/right), bold/italic/underline, letter spacing, line height, fill, shadow. Inline editing via double-click.'],
            ['rect', 'Store default shape',
             'Rectangle with corner radius, fill, stroke color/width, opacity, shadow, full transforms.'],
            ['circle', 'Store / templates',
             'Circle sized by width/height, fill/stroke/opacity and animation.'],
            ['image', 'Toolbar to Image, Assets panel, templates',
             'Raster image from upload or data URL; position, size, rotation, visibility, animation.'],
            ['video', 'Assets panel (video uploads)',
             'Video element (default 320x180) inserted from uploaded video assets; export wires autoplay events.'],
            ['shape', 'Templates / store',
             'SVG Path-based shape with path data, fill/stroke, scalable via width/height.'],
            ['isiScroll', 'Toolbar to ISI, templates',
             'Scrollable compliance ISI tray: isiText (HTML), scroll speed/duration/delay, auto-start, header bar, logo, padding/margin, scrollbar and typography settings (see Section 12).'],
            ['html', 'Canvas/store',
             'Custom HTML content element (htmlContent) wrapped for rendering; edited in the Properties panel HTML Content section.'],
        ],
        widths=[0.8, 1.5, 4.2],
    )
    doc.add_heading('Shared Element Properties', level=2)
    add_bullets(doc, [
        'id, type, name (custom layer name), visible, locked.',
        'x, y, width, height, rotation, scaleX/scaleY, opacity.',
        'anim (manual keyframes), animation/animationDuration/animationDelay/animationLoop (preset).',
        'enterAnimation/enterDelay and exitAnimation/exitDelay (entrance/exit presets).',
        'animations[] — additional timed preset blocks on the global timeline.',
        'hoverAnimation (none | colorChange | glow | shadowPop | letterSpacing | scale) and hoverColor.',
        'Wave properties (waveAmplitude, waveFrequency, waveSpeed, waveColor, wavePoints, waveLayers).',
    ])

    # =====================================================================
    # 7. Canvas
    # =====================================================================
    doc.add_heading('7. Canvas (Design Surface)', level=1)
    doc.add_paragraph(
        'The canvas is a Konva stage (react-konva) rendering the active artboard. It supports '
        'single-artboard editing and a side-by-side multi-artboard view.'
    )
    doc.add_heading('Canvas Features', level=2)
    add_table(
        doc,
        ['Feature', 'Usage', 'Functionality'],
        [
            ['Selection & transform', 'Click an element; drag handles',
             'Konva Transformer for resize/rotate; selected element drives the Properties panel.'],
            ['Drag to move', 'Drag element', 'Updates x/y in the store (with history snapshot).'],
            ['Inline text editing', 'Double-click a text element',
             'DOM textarea overlay; Enter/commit on blur, Escape cancels.'],
            ['ISI editing dialog', 'Double-click an ISI element',
             'Modal with monospace HTML editor for isiText; Done commits, Escape cancels.'],
            ['Keyboard nudge', 'Arrow keys with element selected',
             'Arrow keys move the selected element (skipped while typing in inputs or during playback).'],
            ['Delete', 'Delete / Backspace',
             'Removes the selected element (ignored while focus is in inputs/textareas).'],
            ['Single / All sizes toggle', 'Segmented control on canvas',
             '"Single" edits one size at a time; "All sizes (n)" shows every artboard together.'],
            ['Multi-artboard rendering', 'Enable All sizes',
             'Each artboard renders with its own dimensions and label; stages are registered for export.'],
            ['Playback rendering', 'Preview running',
             'GSAP timeline drives element nodes; ISI auto-scroll runs on its own clock and pauses the global playhead on hover (previewPaused).'],
        ],
        widths=[1.3, 1.7, 3.5],
    )

    # =====================================================================
    # 8. Properties panel
    # =====================================================================
    doc.add_heading('8. Properties Panel (Right Side)', level=1)
    doc.add_paragraph('The right panel has three states: Canvas Settings (nothing selected), '
                      'Element Properties (an element is selected), and Keyframe Editor (when a '
                      'timeline keyframe diamond is selected).')

    doc.add_heading('8.1 Canvas Settings (no selection)', level=2)
    add_table(
        doc,
        ['Section', 'Controls', 'Functionality'],
        [
            ['Background', 'Color picker, Image URL input',
             'Sets canvas background color and optional background image URL.'],
            ['Canvas Size', 'Width / Height (read-only display)',
             'Shows current artboard dimensions in px (change via artboard presets or setSize).'],
            ['Artboards (Sizes)', 'List with activate/add/remove',
             'Select the active artboard, add a new size, or remove one (cannot remove the last).'],
            ['Duration & Loop', 'Total duration input, Loop toggle',
             'Sets master timeline length (min 0.5 s) and whether playback loops.'],
        ],
        widths=[1.3, 2.0, 3.2],
    )

    doc.add_heading('8.2 Keyframe Editor (timeline diamond selected)', level=2)
    doc.add_paragraph('Selecting a keyframe diamond on the timeline replaces the panel content '
                      'with per-keyframe fields:')
    add_bullets(doc, [
        'Time (absolute seconds on the global timeline).',
        'X, Y position; Opacity (0-100); Rotation; Scale X / Scale Y.',
        'Easing — dropdown of preset curves (Linear, Ease In/Out, power1-3, sine, back, elastic, bounce, expo).',
        'Remove keyframe action.',
        'Guidance text when no keyframe is selected: double-click a layer row on the timeline to add one.',
    ])

    doc.add_heading('8.3 Element Properties (element selected)', level=2)
    add_table(
        doc,
        ['Section', 'Applies To', 'Controls / Functionality'],
        [
            ['Element header', 'All', 'Element type badge, arrange buttons (bring to front / forward / backward / send to back).'],
            ['Layer Name', 'All', 'Rename field; name appears in Layers panel and Timeline.'],
            ['Position & Size', 'All', 'X, Y, Width, Height, Rotation (degrees) numeric inputs.'],
            ['Appearance', 'rect, text, circle', 'Fill color picker, corner radius (rect), opacity slider (0-100%), stroke color, stroke width.'],
            ['Typography', 'text', 'Text content textarea, font size, font family, alignment (left/center/right), style toggles (Bold, Italic, Underline), letter spacing, line height, fill color.'],
            ['Shadow', 'rect, text', 'Shadow color, blur, offset X/Y (plus shadow opacity via store fields).'],
            ['Animation', 'All', 'Shows current preset/animation state; "Browse all presets" and "Sequence & Stagger All Layers" buttons open the Animation Studio (see Section 11).'],
            ['HTML Content', 'html', 'Monospace textarea for raw HTML (placeholder "Your HTML here...").'],
            ['ISI Pro Suite', 'isiScroll', 'Full ISI configuration — see Section 12.'],
            ['Image / Video props', 'image, video', 'Source (src), size and transform controls.'],
        ],
        widths=[1.2, 0.9, 4.4],
    )

    # =====================================================================
    # 9. Left panels
    # =====================================================================
    doc.add_heading('9. Left Panels: Layers, Assets, Templates, Sizes', level=1)

    doc.add_heading('9.1 Layers Panel', level=2)
    doc.add_paragraph('Lists all elements front-most first (top of list = on top of canvas). '
                      'Click a row to select; the header shows the layer count and rename hint.')
    add_table(
        doc,
        ['Row Control', 'Usage', 'Functionality'],
        [
            ['Type icon', '—', 'Icon by type: text/rect/circle/image/ISI/shape/HTML.'],
            ['Name', 'Double-click name or click pencil',
             'Inline rename input; Enter/blur commits, Escape cancels. Falls back to "Type XXXX" (last 4 id chars).'],
            ['Bring forward / Send backward', 'Arrow up / Arrow down buttons',
             'Moves the element one step in z-order (reorderElement up/down).'],
            ['Visibility', 'Eye / Eye-off', 'Toggles visible flag; hidden layers render at 40% opacity in the list.'],
            ['Lock', 'Lock / Unlock', 'Toggles locked flag (locked shown in red).'],
            ['Duplicate', 'Copy icon', 'Clones element offset by +12 px with fresh ids for keyframes/blocks.'],
            ['Delete', 'Trash icon', 'Removes the element.'],
        ],
        widths=[1.4, 1.6, 3.5],
    )

    doc.add_heading('9.2 Assets Panel', level=2)
    add_bullets(doc, [
        'Upload button — accepts image and video files (PNG, JPG, WebP, GIF, MP4, WebM); invalid types are rejected with an alert.',
        'Search field — filters uploaded assets by file name.',
        'Persistence — assets are stored in localStorage under "banner_tool_uploaded_assets" as data URLs, so they survive reloads.',
        'Insert — click a thumbnail to add it to the canvas (images default 200x150, videos 320x180 at 100,100).',
        'Delete — hover a tile and click the X to remove it from the library.',
        'File counter — shows how many files are uploaded.',
    ])

    doc.add_heading('9.3 Templates Panel (Template Gallery)', level=2)
    add_bullets(doc, [
        'Category chips (All / Static / Medical / Animated) filter the gallery.',
        'Each template card shows a proportional thumbnail with an ISI scroll area strip, dimensions, duration, and Static/Animated badge (animated cards show an ANIM film badge).',
        'Clicking a card calls loadTemplate(elements, width, height, totalDuration), replacing the current design with the template and resetting history.',
        'Built-in templates: EMR Static Banner (300x250), Master EMR Animated Sequence (300x250, 12 s), EMR Leaderboard (728x90), EMR Large Rectangle (336x280), EMR Mobile Animated (320x100, 8 s).',
    ])

    doc.add_heading('9.4 Sizes Panel (Multi-size Preview / Variations)', level=2)
    add_bullets(doc, [
        'Renders a live Konva thumbnail for every artboard using smart reflow (reflowElements) from the base size.',
        'Each preview shows the label, dimensions, and an active red highlight.',
        'Click a preview to make that artboard active — the canvas switches to it.',
    ])

    # =====================================================================
    # 10. Timeline
    # =====================================================================
    doc.add_heading('10. Timeline (Bottom Panel)', level=1)
    doc.add_paragraph('The timeline is the animation workspace: it shows one row per layer, '
                      'animation segments, keyframe diamonds, and the red playhead.')

    doc.add_heading('10.1 Top Controls', level=2)
    add_bullets(doc, [
        'Playback: play/pause, stop, go-to-start/end buttons tied to the GSAP preview.',
        'Time display: current playhead time / total duration.',
        'Total duration and loop controls (also mirrored in canvas settings).',
        '"Stagger All" button — opens the Animation Studio Layer Sequencer to stagger all layers.',
        'Zoom / "Fit to Timeline Window" (maximize icon) — scales the visible time range.',
        'Snap toggle — "Snap: 0.1s" vs "Smooth Scrub" display in the footer.',
    ])

    doc.add_heading('10.2 Layer Rows', level=2)
    add_table(
        doc,
        ['Control', 'Usage', 'Functionality'],
        [
            ['Drag handle (grip)', 'Drag row vertically', 'Reorders layers (lower in list = in front on canvas).'],
            ['Visibility / Lock', 'Eye and lock icons', 'Same toggles as the Layers panel.'],
            ['Sparkles button', 'Click', 'Selects the layer and opens the Animation Studio on the relevant tab (in/out/loop).'],
            ['Delete', 'Trash icon', 'Removes the layer.'],
            ['Animation segments', 'Drag body to move; drag right edge to resize',
             'Shifts an animation block in time or changes its duration. Double-click opens the Studio or adds a keyframe.'],
            ['Keyframe diamonds', 'Double-click row to add; click diamond to select',
             'Adds/edits manual keyframes; selecting one opens the keyframe editor in the Properties panel.'],
        ],
        widths=[1.3, 1.8, 3.4],
    )

    doc.add_heading('10.3 Playhead & Footer Shortcuts', level=2)
    doc.add_paragraph('Drag the red diamond/line to scrub the playhead. The footer lists the '
                      'built-in shortcuts:')
    add_table(
        doc,
        ['Shortcut', 'Action'],
        [
            ['Space', 'Play / Pause'],
            ['Left / Right arrows', 'Step 0.1 s'],
            ['Drag red line', 'Scrub playhead'],
            ['Drag segment', 'Adjust timing'],
            ['Double click', 'Open Studio / add keyframe'],
            ['Footer mode', '"Snap: 0.1s" or "Smooth Scrub" indicator'],
        ],
        widths=[1.8, 4.7],
    )

    # =====================================================================
    # 11. Animation Studio
    # =====================================================================
    doc.add_heading('11. Animation Studio (Animation Dialog)', level=1)
    doc.add_paragraph(
        'A full-screen modal opened from the Properties panel ("Browse all presets" / '
        '"Sequence & Stagger All Layers"), the Timeline ("Stagger All", sparkles buttons), or '
        'programmatically via the "open-animation-studio" window event (with a detail.tab hint).'
    )

    doc.add_heading('11.1 Tabs', level=2)
    add_table(
        doc,
        ['Tab', 'Color Dot', 'Purpose'],
        [
            ['Entrance (In)', 'Amber', 'Preset entrance animations applied as enterAnimation (or animation) on the selected layer.'],
            ['Emphasis & Loop', 'Purple', 'Looping/emphasis presets; sets animation with animationLoop where flagged.'],
            ['Exit (Out)', 'Red', 'Exit presets applied as exitAnimation.'],
            ['Layer Sequencer', 'Emerald', 'Batch stagger timing across all layers.'],
        ],
        widths=[1.4, 0.9, 4.2],
    )
    doc.add_paragraph('The header shows the selected layer chip (or "Select a layer on canvas" '
                      'warning) and a close button.')

    doc.add_heading('11.2 Preset Browser (Entrance / Loop / Exit tabs)', level=2)
    add_bullets(doc, [
        'Search box — full-text search over preset names/ids (e.g. "fade, slide, bounce, pop, 3d, blur").',
        'Category chips — filter by family: Fade, Slide, Scale & Zoom, 3D Flip, Rotate & Roll, Bounce & Elastic, Puff & Slit, Text & Typography, Glitch & Blur, Emphasis & Loop, Ken Burns, etc.',
        'Preset grid — cards from the Animista-inspired catalog (60+ register() calls expanding to directional variants, 300+ animations overall), grouped into Entrance, Loop/Emphasis, and Exit families.',
        'Direction matrix — 9-way position picker (Top-Left, Top, Top-Right, Left, Center, Right, Bottom-Left, Bottom, Bottom-Right) for directional presets.',
        'Timing controls — Duration (seconds, default 0.8) and Delay (entrance delay for "in"; gap before exit for "out"; generic delay otherwise), plus an easing selector.',
        'Apply — writes the preset id + timing to the selected element (setting a preset clears conflicting manual keyframes to avoid duplicate animations).',
    ])
    doc.add_heading('Animation families in the catalog', level=3)
    add_bullets(doc, [
        'Entrance (IN): Fade In (9 directions), Slide In (8 directions plus blurred variants), Scale In / Pop, 3D Flip In, Rotate & Roll In, Swirl In, Bounce & Jello In, Puff & Slit In, Swing In, Text & Typography (tracking in, focus, pop-up, neon flicker), Glitch & Flicker (strobe, camera blur, digital glitch).',
        'Loop & Emphasis: Heartbeat Pulse, Ken Burns & background loops, and other attention styles (auto-loop flagged presets).',
        'Exit (OUT): Fade Out, Slide Out, Scale Out & Puff, 3D Flip Out, Rotate & Roll Out.',
        'Hover effects (element property): none, colorChange, glow, shadowPop, letterSpacing, scale.',
    ])

    doc.add_heading('11.3 Layer Sequencer (Stagger)', level=2)
    add_table(
        doc,
        ['Setting', 'Options', 'Functionality'],
        [
            ['Gap', 'Numeric seconds (default 0.2)', 'Delay inserted between consecutive layers.'],
            ['Order', 'Bottom to Top / Top to Bottom / Center Outward / Random Shuffle',
             'Sorting applied before assigning incremental delays (applyStaggeredDelays).'],
            ['Target', 'Entrance / Main', 'Whether stagger applies to entrance delays or the main/timed animation.'],
            ['Apply', 'Stagger button', 'Batch-updates delays across all layers in one action.'],
        ],
        widths=[0.9, 2.4, 3.2],
    )

    # =====================================================================
    # 12. ISI
    # =====================================================================
    doc.add_heading('12. ISI Pro Suite (Compliance Scroll)', level=1)
    doc.add_paragraph(
        'The isiScroll element implements the pharma-required Important Safety Information '
        'tray. It is added from the toolbar (ISI), from templates, or via HTML export CSS '
        'classes. The ISI scroll timeline is independent of the master banner timeline — it '
        'runs on its own clock, auto-starts, and hovering during preview pauses the global '
        'playhead (previewPaused).'
    )
    doc.add_heading('Properties (ISI Pro Suite section in Properties panel)', level=2)
    add_table(
        doc,
        ['Group', 'Properties', 'Functionality'],
        [
            ['Content', 'ISI Text (HTML supported)',
             'Rich HTML safety content; toolbar buttons wrap selections, insert a Prescribing Information (PI) link bar, or insert a bullet list.'],
            ['Header Bar (PI Link Strip)', 'isiHeaderText, isiHeaderLink, isiHeaderColor, isiHeaderBackground, isiHeaderHeight',
             'Traditional "Prescribing Information" patient link strip above the ISI body.'],
            ['Logo / Branding', 'isiLogoSrc, isiLogoLink, isiLogoWidth, isiLogoPosition (top/bottom)',
             'Optional logo image with link inside the ISI tray.'],
            ['Styling & Scroll', 'isiScrollSpeed, isiAutoStart, isiStartDelay, isiScrollDuration, isiBackgroundColor, isiBorderWidth/Color, isiPadding* and isiMargin*',
             'Controls auto-scroll behavior (start delay on the ISI clock, full scroll duration before reset, default 60 s) and tray chrome.'],
            ['Scrollbar', 'isiScrollbarColor, isiScrollbarTrackColor, isiScrollbarHeight/Width, isiScrollbarMarginTop/MarginRight, isiScrollbarPadding',
             'Custom scrollbar appearance and spacing.'],
            ['ISI Typography', 'isiFontFamily, isiFontWeight, isiFontStyle, isiLineHeight, isiLetterSpacing, plus fontSize/fill',
             'Type styling for the safety content.'],
        ],
        widths=[1.2, 2.4, 3.1],
    )
    doc.add_heading('Runtime Behavior', level=2)
    add_bullets(doc, [
        'Auto-scroll begins after isiStartDelay (0 = immediately) on the ISI tray own clock.',
        'Content scrolls top to bottom over isiScrollDuration seconds, then resets and repeats.',
        'Mouse hover over ISI content pauses the global playhead until the mouse leaves.',
        'Inline ISI editing: double-click the element to open the HTML editor modal.',
        'Exported packages include full ISI CSS (isi_wrapper, scrollbar classes, patient_link PI strip) for drop-in parity.',
    ])

    # =====================================================================
    # 13. Artboards
    # =====================================================================
    doc.add_heading('13. Artboards & Multi-size Campaigns', level=1)
    doc.add_paragraph('Each artboard stores its own element list, dimensions, and label. The '
                      'store keeps the active artboard elements in sync (persistActiveElements).')
    doc.add_heading('Standard IAB Size Presets', level=2)
    add_table(
        doc,
        ['Preset', 'Width x Height'],
        [
            ['Medium Rectangle', '300 x 250'],
            ['Leaderboard', '728 x 90'],
            ['Wide Skyscraper', '160 x 600'],
            ['Half Page', '300 x 600'],
            ['Billboard', '970 x 250'],
            ['Large Rectangle', '336 x 280'],
            ['Square', '250 x 250'],
            ['Mobile Leaderboard', '320 x 50'],
            ['Mobile Banner', '320 x 100'],
            ['Skyscraper', '120 x 600'],
        ],
        widths=[3.0, 3.0],
    )
    doc.add_heading('Multi-size Campaign Workflow', level=2)
    add_numbered(doc, [
        'Click "Multi-size" in the header to open the campaign modal.',
        'Use "Select all" or tick individual Standard IAB sizes (already-added sizes are disabled and marked "added").',
        'Click "Generate n sizes" — addCampaignSizes creates the selected artboards from the current design.',
        'Switch between sizes via header tabs or the Sizes panel; use "All sizes" on the canvas to edit them together.',
        'setSize applies smart reflow (reflowElements) so layouts scale proportionally when dimensions change.',
    ])

    # =====================================================================
    # 14. Export
    # =====================================================================
    doc.add_heading('14. Export Functionality', level=1)
    add_table(
        doc,
        ['Export', 'Trigger', 'Output & Details'],
        [
            ['PNG image', 'Header "PNG" button (export-canvas event)',
             'Deselects elements, captures the active artboard Konva stage at pixelRatio 2, downloads banner-WxH.png.'],
            ['HTML package', 'Header "Export HTML" button (export-html event)',
             'JSZip archive named banner-package-WxH.zip containing: index.html; css/styles.css (banner box, full ISI stylesheet, @keyframes with cubic-bezier timing functions mapped from GSAP eases); js/main.js (timeline built from unit-tested builders); images/ (collected assets).'],
        ],
        widths=[1.1, 1.9, 3.7],
    )
    doc.add_heading('Export Notes', level=2)
    add_bullets(doc, [
        'GSAP eases are mapped to CSS timing functions (power1-4, sine, expo, back, elastic, bounce, stepped); unmapped eases fall back to linear.',
        'Keyframe times become percentages of the master duration; elements can be hidden in CSS until their animation starts.',
        'Video elements emit autoplay wiring in the generated main.js.',
        'The emitted JavaScript is kept syntactically valid by design (pure builders in bannerExport.ts, covered by unit tests) so the banner always animates when opened.',
        'Loop and total duration from the editor are honored in the package.',
    ])

    # =====================================================================
    # 15. Templates
    # =====================================================================
    doc.add_heading('15. Templates Catalog', level=1)
    add_table(
        doc,
        ['Template', 'Mode', 'Category', 'Size', 'Duration'],
        [
            ['EMR Static Banner', 'Static (emr)', 'Static / Medical', '300x250', '6 s'],
            ['Master EMR Animated Sequence', 'Animated', 'Animated / Medical', '300x250', '12 s'],
            ['EMR Leaderboard', 'Static (emr)', 'Static / Medical', '728x90', '6 s'],
            ['EMR Large Rectangle', 'Static (emr)', 'Static / Medical', '336x280', '6 s'],
            ['EMR Mobile Animated', 'Animated', 'Animated / Medical', '320x100', '8 s'],
        ],
        widths=[2.1, 1.0, 1.4, 0.9, 0.7],
    )
    doc.add_paragraph(
        'All templates embed the pre-loaded EMR_ISI_TEXT (ORSERDU Important Safety Info & '
        'Indication HTML) in an isiScroll element. The Master EMR Animated Sequence '
        'demonstrates keyframe choreography: three headline messages fade/slide in sequence, a '
        'CTA pops in with back.out easing, and a background orb/logo animate on loop.'
    )

    # =====================================================================
    # 16. Shortcuts
    # =====================================================================
    doc.add_heading('16. Keyboard Shortcuts', level=1)
    add_table(
        doc,
        ['Shortcut', 'Context', 'Action'],
        [
            ['Ctrl+Z', 'Global', 'Undo'],
            ['Ctrl+Shift+Z', 'Global', 'Redo'],
            ['Ctrl+Y', 'Global', 'Redo'],
            ['Delete / Backspace', 'Canvas, element selected, not typing in a field', 'Delete selected element'],
            ['Arrow keys', 'Canvas, element selected, not playing', 'Nudge element by keyboard step'],
            ['Space', 'Timeline focused', 'Play / Pause'],
            ['Left / Right arrows', 'Timeline focused', 'Step playhead 0.1 s'],
            ['Double-click', 'Timeline row / segment', 'Open Animation Studio or add keyframe'],
            ['Double-click', 'Canvas text element', 'Inline text edit'],
            ['Double-click', 'Canvas ISI element', 'Open ISI HTML editor'],
            ['Double-click', 'Layer name (Layers panel)', 'Rename layer'],
            ['Enter / Escape', 'Rename inputs, editors', 'Commit / cancel edit'],
        ],
        widths=[1.4, 2.6, 2.5],
    )

    # =====================================================================
    # 17. State management
    # =====================================================================
    doc.add_heading('17. State Management & Undo/Redo', level=1)
    doc.add_paragraph('All UI regions read/write a single Zustand store (designStore.ts). '
                      'Mutating actions push a history snapshot (saveHistory) capturing elements, '
                      'canvas size, duration, loop, background, artboards, and the active artboard.')
    doc.add_heading('Store Actions by Group', level=2)
    add_table(
        doc,
        ['Group', 'Actions'],
        [
            ['Elements', 'addElement, updateElement, duplicateElement, selectElement, removeElement, reorderElement (up/down/top/bottom), moveElement (exact z-index), toggleVisibility, toggleLock'],
            ['Keyframes', 'addKeyframe, updateKeyframe, removeKeyframe, selectKeyframe'],
            ['Playback', 'setPlayheadTime, setIsPlaying, setPreviewPaused, setTotalDuration (min 0.5 s), setLoop'],
            ['Artboards', 'addArtboard, addCampaignSizes, removeArtboard, setActiveArtboard, setMultiArtboardView'],
            ['Global', 'reset, setSize (with smart reflow), setCanvasBackground, setCanvasBackgroundImage, loadTemplate'],
            ['History', 'undo, redo, clearHistory'],
        ],
        widths=[1.2, 5.3],
    )
    doc.add_heading('Conflict Rules', level=2)
    add_bullets(doc, [
        'Setting any preset animation field (enterAnimation, exitAnimation, animations, animation, durations, delays) clears manual anim keyframes to prevent duplicate/conflicting animations.',
        'Duplicating an element clones and re-ids its keyframes and timed animation blocks.',
        'History excludes transient state (selection, playhead, playing flags).',
    ])

    # =====================================================================
    # 18. Tech stack
    # =====================================================================
    doc.add_heading('18. Technology Stack', level=1)
    add_table(
        doc,
        ['Layer', 'Technology', 'Purpose'],
        [
            ['UI framework', 'React 19.2', 'Component-based interface'],
            ['Language', 'TypeScript 5.9', 'Type safety across store/components/utils'],
            ['State', 'Zustand 5.0', 'Global store with undo/redo history'],
            ['Canvas', 'Konva 10 + react-konva 19.2', '2D stage rendering and transforms'],
            ['Animation', 'GSAP 3.14', 'Preview timeline playback (AnimationHelpers)'],
            ['Styling', 'Tailwind CSS 3.4 + PostCSS', 'Utility-first UI styling'],
            ['Icons', 'Lucide React', 'Toolbar/panel iconography'],
            ['Export', 'JSZip 3.10', 'HTML/CSS/JS/images ZIP packaging'],
            ['Images', 'use-image', 'Async image loading for Konva'],
            ['Build', 'Vite 7.2', 'Dev server and production bundling'],
            ['Testing', 'Vitest 4.1', 'Unit tests (store, keyframes, bannerExport, animation helpers)'],
            ['Linting', 'ESLint 9.3', 'Code quality'],
            ['Utilities', 'clsx, tailwind-merge', 'Class-name helpers'],
        ],
        widths=[1.2, 2.0, 3.3],
    )

    doc.add_heading('Commands', level=2)
    add_bullets(doc, [
        'npm run dev — start the Vite dev server.',
        'npm run build — type-check (tsc -b) and build to dist/.',
        'npm run preview — preview the production build.',
        'npm run test — run Vitest suites once (test:watch for watch mode, coverage for reports).',
        'npm run lint — run ESLint.',
    ])

    doc.save(OUTPUT)
    print(f'Created: {OUTPUT}')


if __name__ == '__main__':
    build()

