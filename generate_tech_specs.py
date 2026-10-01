#!/usr/bin/env python3
"""
Generate the Tech Specifications document in Word format
(Banner Tool - Animated_banner_tool)
"""

import os

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def shade_cell(cell, color):
    """Add shading to a cell"""
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:fill'), color)
    cell._element.get_or_add_tcPr().append(shading_elm)


def add_table(doc, headers, rows, header_color='2E74B5'):
    """Add a formatted table with a shaded header row"""
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        for paragraph in hdr_cells[i].paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.size = Pt(10)
        shade_cell(hdr_cells[i], header_color)
    for row_data in rows:
        row_cells = table.add_row().cells
        for i, value in enumerate(row_data):
            row_cells[i].text = str(value)
            for paragraph in row_cells[i].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
    return table


def add_code_block(doc, text):
    """Add a monospaced code block paragraph"""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:fill'), 'F2F2F2')
    p._element.get_or_add_pPr().append(shading_elm)
    return p

def create_word_doc():
    """Create the Tech Specifications Word document"""
    doc = Document()

    title = doc.add_heading('Banner Tool - Technical Specifications', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph('Animated Banner Tool - Technology Stack & Engineering Specifications')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(12)
    subtitle.runs[0].font.italic = True

    doc.add_paragraph('Generated: 2026-09-30')
    doc.add_paragraph()

    doc.add_heading('Table of Contents', 1)
    toc_items = [
        '1. Project Overview',
        '2. Core Technology Stack',
        '3. Utility Libraries',
        '4. Testing & Code Quality',
        '5. TypeScript Configuration',
        '6. NPM Scripts & Commands',
        '7. System Architecture',
        '8. Source Code Structure',
        '9. Environment & Configuration',
        '10. Firebase Backend Services',
    ]
    for item in toc_items:
        doc.add_paragraph(item, style='List Bullet')

    doc.add_page_break()

    # 1. Project Overview
    doc.add_heading('1. Project Overview', 1)
    doc.add_paragraph(
        'The Banner Tool is a professional React-based animated banner designer built for creating '
        'compliant pharmaceutical and healthcare marketing materials (EMR and animated banner modes). '
        'It provides a canvas-based editing experience with keyframe animation, templates, and '
        'Firebase-backed project persistence.'
    )
    add_table(doc, ['Attribute', 'Value'], [
        ['Package name', 'banner-tool'],
        ['Version', '0.0.0 (private)'],
        ['Module system', 'ESM ("type": "module")'],
        ['Platform', 'Web (browser SPA)'],
        ['Repository', 'https://github.com/darshant419/Animated_banner_tool.git'],
        ['Branch / Commit', 'main / ce891c068c9ea213b0193db9735d5b7060b40816'],
    ])
    doc.add_paragraph()

    # 2. Core Technology Stack
    doc.add_heading('2. Core Technology Stack', 1)
    add_table(doc, ['Layer', 'Technology', 'Version'], [
        ['UI Framework', 'React + React DOM', '^19.2.0'],
        ['Language', 'TypeScript (strict mode)', '~5.9.3'],
        ['Build Tool', 'Vite (+ @vitejs/plugin-react)', '^7.2.4'],
        ['Styling', 'Tailwind CSS + PostCSS + Autoprefixer', 'tailwindcss ^3.4.17'],
        ['State Management', 'Zustand', '^5.0.9'],
        ['Canvas Rendering', 'Konva + react-konva (+ use-image)', 'konva ^10.0.12 / react-konva ^19.2.1'],
        ['Animation Engine', 'GSAP', '^3.14.2'],
        ['Backend / Services', 'Firebase (Auth, Firestore, Storage)', '^12.19.0'],
        ['Routing', 'Custom hash router (no react-router)', 'src/router/hashRouter.ts'],
        ['Icons', 'lucide-react', '^0.562.0'],
    ])
    doc.add_paragraph()

    # 3. Utility Libraries
    doc.add_heading('3. Utility Libraries', 1)
    add_table(doc, ['Library', 'Version', 'Purpose'], [
        ['jszip', '^3.10.1', 'Export / packaging of banners'],
        ['lucide-react', '^0.562.0', 'Icon set'],
        ['clsx', '^2.1.1', 'Conditional class name utilities'],
        ['tailwind-merge', '^3.4.0', 'Tailwind class deduplication / merging'],
        ['use-image', '^1.1.4', 'Konva image loading hook'],
    ])
    doc.add_paragraph()
    # 4. Testing & Code Quality
    doc.add_heading('4. Testing & Code Quality', 1)
    add_table(doc, ['Tool', 'Version', 'Configuration'], [
        ['Vitest', '^4.1.11', "environment: 'node'; include: src/**/*.test.ts"],
        ['@vitest/coverage-v8', '^4.1.11', "reporter: ['text', 'html']"],
        ['ESLint', '^9.39.1', 'Flat config (eslint.config.js)'],
        ['typescript-eslint', '^8.46.4', 'Type-aware lint rules'],
        ['eslint-plugin-react-hooks', '^7.0.1', 'Hooks rules'],
        ['eslint-plugin-react-refresh', '^0.4.24', 'HMR refresh rules'],
    ])
    doc.add_paragraph()
    doc.add_paragraph(
        'Test files are colocated with source code (e.g. designStore.test.ts, assetService.test.ts). '
        'The src directory contains 30 TypeScript and 26 TSX source files, plus 2 CSS files.'
    )
    doc.add_paragraph()

    # 5. TypeScript Configuration
    doc.add_heading('5. TypeScript Configuration', 1)
    add_table(doc, ['Option', 'Value'], [
        ['Target', 'ES2022'],
        ['Lib', 'ES2022, DOM, DOM.Iterable'],
        ['Module', 'ESNext'],
        ['Module resolution', 'bundler'],
        ['JSX', 'react-jsx'],
        ['Strict', 'true'],
        ['noUnusedLocals / noUnusedParameters', 'true'],
        ['noEmit', 'true (type-check via tsc -b)'],
        ['Other', 'verbatimModuleSyntax, erasableSyntaxOnly, noFallthroughCasesInSwitch, noUncheckedSideEffectImports'],
    ])
    doc.add_paragraph()
    doc.add_paragraph('Project references: tsconfig.json -> tsconfig.app.json (src) + tsconfig.node.json (vite config).')
    doc.add_paragraph()

    # 6. NPM Scripts
    doc.add_heading('6. NPM Scripts & Commands', 1)
    add_code_block(doc,
        'dev       -> vite                    # Start dev server (host: true, LAN accessible)\n'
        'build     -> tsc -b && vite build    # Type-check then production build\n'
        'lint      -> eslint .                # Lint entire project\n'
        'test      -> vitest run              # Run tests once\n'
        'test:watch-> vitest                  # Watch mode\n'
        'coverage  -> vitest run --coverage   # Coverage report (V8)\n'
        'preview   -> vite preview            # Preview production build'
    )
    doc.add_paragraph()
    # 7. System Architecture
    doc.add_heading('7. System Architecture', 1)
    doc.add_paragraph(
        'High-level flow (per ARCHITECTURE.md): main.tsx -> App.tsx -> ModeSelect (EMR vs Animated) '
        '-> MainLayout, which orchestrates Toolbar, DesignCanvas (Konva), PropertiesPanel, Timeline, '
        'LayersPanel, AssetsPanel, TemplatesPanel and VariationsPanel - all wired to a central '
        'Zustand design store with undo/redo history.'
    )
    add_table(doc, ['Subsystem', 'Key Modules'], [
        ['Canvas rendering', 'ElementShape.tsx, ISIScroll.tsx, ISIOverlay.tsx, Konva primitives (Stage, Layer, Rect, Circle, Text, Image, Path, Transformer)'],
        ['Animation & effects', 'AnimationHelpers.ts (GSAP), keyframes.ts, animations.ts, Animista catalog (300+ presets)'],
        ['Data models', 'DesignElement, Artboard, AnimationKeyframe, ElementAnimation, ElementTimedAnimation'],
        ['Template system', 'emrTemplates.ts, BannerTemplate (EMR + Animated), EMR ISI text'],
        ['Utility layer', 'use-image, jszip, lucide-react, Tailwind CSS, clsx'],
        ['Build tools', 'Vite, TypeScript, Vitest, ESLint'],
    ])
    doc.add_paragraph()

    # 8. Source Code Structure
    doc.add_heading('8. Source Code Structure', 1)
    add_code_block(doc,
        'src/\n'
        '  assets/        components/   (AnimationDialog, AssetsPanel, Canvas, LayersPanel,\n'
        '                                Layout, ModeSelect, Projects, PropertiesPanel, Shell,\n'
        '                                Templates, TemplatesPanel, Timeline, Toolbar, Variations)\n'
        '  pages/         router/       (hashRouter.ts, useHashRoute.ts + tests)\n'
        '  services/      (assetService, projectService, templateService, storageService,\n'
        '                   firebase.ts, firestoreUtils.ts + tests)\n'
        '  store/         (designStore.ts + test)\n'
        '  templates/     utils/'
    )
    doc.add_paragraph()

    # 9. Environment & Configuration
    doc.add_heading('9. Environment & Configuration', 1)
    doc.add_paragraph('Firebase configuration is provided via .env (see .env.example) with VITE_ prefixed variables:')
    add_table(doc, ['Variable', 'Description'], [
        ['VITE_FIREBASE_API_KEY', 'Firebase API key'],
        ['VITE_FIREBASE_AUTH_DOMAIN', 'Auth domain (project.firebaseapp.com)'],
        ['VITE_FIREBASE_PROJECT_ID', 'Firebase project ID'],
        ['VITE_FIREBASE_STORAGE_BUCKET', 'Storage bucket'],
        ['VITE_FIREBASE_MESSAGING_SENDER_ID', 'Cloud Messaging sender ID'],
        ['VITE_FIREBASE_APP_ID', 'Firebase app ID'],
    ])
    doc.add_paragraph()
    doc.add_paragraph(
        'Other config files: vite.config.ts (React plugin, server.host = true), vitest.config.ts, '
        'tailwind.config.js (content: index.html + src/**/*.{js,ts,jsx,tsx}), postcss.config.js, '
        'eslint.config.js, cors.json, firestore.rules, storage.rules.'
    )
    doc.add_paragraph()

    # 10. Firebase Backend Services
    doc.add_heading('10. Firebase Backend Services', 1)
    add_table(doc, ['Service', 'Usage'], [
        ['Firestore', 'Project & template persistence (projectService, templateService)'],
        ['Storage', 'Asset storage (storageService, assetService)'],
        ['Security rules', 'firestore.rules, storage.rules'],
        ['CORS', 'cors.json (storage bucket CORS configuration)'],
    ])
    doc.add_paragraph()
    doc.add_paragraph(
        'Note: The README.md is still the default Vite React template readme; the authoritative '
        'architecture reference is ARCHITECTURE.md and the Banner_Tool_Architecture documents.'
    )

    output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'Tech_Specifications_Banner_Tool.docx')
    doc.save(output_path)
    print(f'Word document created: {output_path}')


if __name__ == '__main__':
    create_word_doc()



